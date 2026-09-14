import openpyxl
import pulp
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
def load_data(filepath):
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb['Sheet1']
    rows = list(ws.iter_rows(values_only=True))
    data = rows[1:]
    T = len(data)
    dt = 10 / 60
    times_str = []
    prices, loads, pvs = [], [], []
    for r in data:
        t = r[0]
        if isinstance(t, str):
            if '+1' in t or '0:00+1' in t:
                times_str.append("24:00")
            else:
                times_str.append(t)
        else:
            times_str.append(f"{t.hour:02d}:{t.minute:02d}")
        prices.append(float(r[1]))
        loads.append(float(r[2]))
        pvs.append(float(r[3]))
    return times_str, prices, loads, pvs, T, dt
E_MAX = 10800.0
E_MIN = 1200.0
E_INIT = 6000.0
P_CH_MAX = 5000.0
P_DIS_MAX = 5000.0
ETA_CH = 0.90
ETA_DIS = 0.90
def solve_microgrid(prices, loads, pvs, T, dt):
    prob = pulp.LpProblem("Microgrid_Problem1", pulp.LpMinimize)
    P_buy = [pulp.LpVariable(f'P_buy_{t}', lowBound=0) for t in range(T)]
    P_ch = [pulp.LpVariable(f'P_ch_{t}', lowBound=0, upBound=P_CH_MAX) for t in range(T)]
    P_dis = [pulp.LpVariable(f'P_dis_{t}', lowBound=0, upBound=P_DIS_MAX) for t in range(T)]
    E = [pulp.LpVariable(f'E_{t}', lowBound=E_MIN, upBound=E_MAX) for t in range(T + 1)]
    prob += pulp.lpSum(prices[t] * P_buy[t] * dt for t in range(T)), "Total_Cost"
    for t in range(T):
        prob += P_buy[t] + pvs[t] + P_dis[t] == loads[t] + P_ch[t], f'Power_Balance_{t}'
    prob += E[0] == E_INIT, 'E_initial'
    for t in range(T):
        prob += E[t+1] == E[t] + ETA_CH * P_ch[t] * dt - P_dis[t] * dt / ETA_DIS, f'SOC_Dynamics_{t}'
    prob += E[T] == E_INIT, 'E_final'
    prob.solve(pulp.PULP_CBC_CMD(msg=0))
    if prob.status != pulp.LpStatusOptimal:
        raise RuntimeError("模型求解失败！请检查约束条件。")
    P_buy_val = [P_buy[t].varValue for t in range(T)]
    P_ch_val = [P_ch[t].varValue for t in range(T)]
    P_dis_val = [P_dis[t].varValue for t in range(T)]
    E_val = [E[t].varValue for t in range(T + 1)]
    return prob, P_buy_val, P_ch_val, P_dis_val, E_val
def verify_results(P_buy, P_ch, P_dis, E, prices, loads, pvs, T, dt):
    balance_ok = True
    for t in range(T):
        lhs = P_buy[t] + pvs[t] + P_dis[t]
        rhs = loads[t] + P_ch[t]
        if abs(lhs - rhs) > 0.01:
            balance_ok = False
    capacity_ok = all(E_MIN - 0.01 <= e <= E_MAX + 0.01 for e in E)
    power_ok = all(0 - 0.01 <= p <= P_CH_MAX + 0.01 for p in P_ch) and \
               all(0 - 0.01 <= p <= P_DIS_MAX + 0.01 for p in P_dis)
    cycle_ok = abs(E[0] - E[T]) < 0.01
    nonneg_ok = all(p >= -0.01 for p in P_buy)
    total_buy = sum(P_buy) * dt
    total_pv = sum(pvs) * dt
    total_load = sum(loads) * dt
    total_ch = sum(P_ch) * dt
    total_dis = sum(P_dis) * dt
    ch_loss = total_ch * (1 - ETA_CH)
    dis_loss = total_dis * (1/ETA_DIS - 1)
    energy_in = total_buy + total_pv
    energy_out = total_load + ch_loss + dis_loss
    energy_balance_ok = abs(energy_in - energy_out) < 1.0
    all_ok = balance_ok and capacity_ok and power_ok and cycle_ok and nonneg_ok and energy_balance_ok
    return all_ok
def generate_tables(times_str, prices, loads, pvs, P_buy, P_ch, P_dis, E, T, dt):
    table1_periods = [
        ("10:00-10:10", "10:10"),
        ("12:00-12:10", "12:10"),
        ("14:00-14:10", "14:10"),
        ("16:00-16:10", "16:10"),
        ("18:00-18:10", "18:10"),
        ("20:00-20:10", "20:10"),
    ]
    table1_data = []
    for period_name, time_tag in table1_periods:
        idx = times_str.index(time_tag)
        buy_kwh = P_buy[idx] * dt
        table1_data.append((period_name, buy_kwh))
    total_buy = sum(P_buy) * dt
    total_cost = sum(prices[t] * P_buy[t] * dt for t in range(T))
    table2_periods = [
        ("0:00-4:00",   0,  24),
        ("4:00-8:00",   24, 48),
        ("8:00-12:00",  48, 72),
        ("12:00-16:00", 72, 96),
        ("16:00-20:00", 96, 120),
        ("20:00-24:00", 120, 144),
    ]
    table2_data = []
    for period_name, start, end in table2_periods:
        ch_kwh = sum(P_ch[t] * dt for t in range(start, end))
        dis_kwh = sum(P_dis[t] * dt for t in range(start, end))
        table2_data.append((period_name, ch_kwh, dis_kwh))
    return {
        'table1': table1_data,
        'table2': table2_data,
        'total_buy': total_buy,
        'total_cost': total_cost,
        'E_0': E[0],
        'E_24': E[T],
    }
def save_to_excel(times_str, prices, loads, pvs, P_buy, P_ch, P_dis, E, T, dt, tables):
    header_font = Font(name='宋体', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    center = Alignment(horizontal='center', vertical='center')
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "表1_购电量"
    ws1['A1'] = "表1 微网在指定时间段的购电量及全天的购电量和购电费"
    ws1.merge_cells('A1:B1')
    ws1['A1'].font = Font(name='宋体', size=12, bold=True)
    ws1['A1'].alignment = center
    headers1 = ['时间段', '购电量(kWh)']
    for col, h in enumerate(headers1, 1):
        c = ws1.cell(row=2, column=col, value=h)
        c.font = header_font
        c.alignment = center
        c.border = thin_border
        c.fill = header_fill
    for i, (period, buy_kwh) in enumerate(tables['table1']):
        ws1.cell(row=3+i, column=1, value=period).alignment = center
        ws1.cell(row=3+i, column=2, value=round(buy_kwh, 4)).alignment = center
        ws1.cell(row=3+i, column=1).border = thin_border
        ws1.cell(row=3+i, column=2).border = thin_border
    row_total = 3 + len(tables['table1'])
    ws1.cell(row=row_total, column=1, value="全天购电量(kWh)").alignment = center
    ws1.cell(row=row_total, column=2, value=round(tables['total_buy'], 4)).alignment = center
    ws1.cell(row=row_total+1, column=1, value="全天购电费(元)").alignment = center
    ws1.cell(row=row_total+1, column=2, value=round(tables['total_cost'], 4)).alignment = center
    for r in [row_total, row_total+1]:
        for c in [1, 2]:
            ws1.cell(row=r, column=c).border = thin_border
            ws1.cell(row=r, column=c).font = Font(name='宋体', size=11, bold=True)
    ws1.column_dimensions['A'].width = 22
    ws1.column_dimensions['B'].width = 20
    ws2 = wb.create_sheet("表2_储能充放电")
    ws2['A1'] = "表2 储能设备在指定时间段的充放电量及0:00和24:00的储电量"
    ws2.merge_cells('A1:C1')
    ws2['A1'].font = Font(name='宋体', size=12, bold=True)
    ws2['A1'].alignment = center
    headers2 = ['时间段', '充电量(kWh)', '放电量(kWh)']
    for col, h in enumerate(headers2, 1):
        c = ws2.cell(row=2, column=col, value=h)
        c.font = header_font
        c.alignment = center
        c.border = thin_border
        c.fill = header_fill
    for i, (period, ch_kwh, dis_kwh) in enumerate(tables['table2']):
        ws2.cell(row=3+i, column=1, value=period).alignment = center
        ws2.cell(row=3+i, column=2, value=round(ch_kwh, 4)).alignment = center
        ws2.cell(row=3+i, column=3, value=round(dis_kwh, 4)).alignment = center
        for c in [1, 2, 3]:
            ws2.cell(row=3+i, column=c).border = thin_border
    row_end = 3 + len(tables['table2'])
    ws2.cell(row=row_end, column=1, value="0:00 储电量(kWh)").alignment = center
    ws2.cell(row=row_end, column=2, value=round(tables['E_0'], 4)).alignment = center
    ws2.cell(row=row_end+1, column=1, value="24:00 储电量(kWh)").alignment = center
    ws2.cell(row=row_end+1, column=2, value=round(tables['E_24'], 4)).alignment = center
    for r in [row_end, row_end+1]:
        for c in [1, 2, 3]:
            ws2.cell(row=r, column=c).border = thin_border
            ws2.cell(row=r, column=c).font = Font(name='宋体', size=11, bold=True)
    ws2.column_dimensions['A'].width = 22
    ws2.column_dimensions['B'].width = 20
    ws2.column_dimensions['C'].width = 20
    ws3 = wb.create_sheet("完整策略")
    headers3 = ['时间段', '电价(元/kWh)', '负载(kW)', '光伏(kW)',
                '购电量(kWh)', '充电量(kWh)', '放电量(kWh)', '储电量(kWh)', '购电费(元)']
    for col, h in enumerate(headers3, 1):
        c = ws3.cell(row=1, column=col, value=h)
        c.font = header_font
        c.alignment = center
        c.border = thin_border
        c.fill = header_fill
    for t in range(T):
        row = 2 + t
        ws3.cell(row=row, column=1, value=times_str[t])
        ws3.cell(row=row, column=2, value=round(prices[t], 4))
        ws3.cell(row=row, column=3, value=round(loads[t], 4))
        ws3.cell(row=row, column=4, value=round(pvs[t], 4))
        ws3.cell(row=row, column=5, value=round(P_buy[t] * dt, 4))
        ws3.cell(row=row, column=6, value=round(P_ch[t] * dt, 4))
        ws3.cell(row=row, column=7, value=round(P_dis[t] * dt, 4))
        ws3.cell(row=row, column=8, value=round(E[t], 4))
        ws3.cell(row=row, column=9, value=round(prices[t] * P_buy[t] * dt, 4))
        for c in range(1, 10):
            ws3.cell(row=row, column=c).alignment = center
            ws3.cell(row=row, column=c).border = thin_border
    for col_letter in ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H', 'I']:
        ws3.column_dimensions[col_letter].width = 16
    output_path = 'result1.xlsx'
    wb.save(output_path)
def calc_no_ess_cost(prices, loads, pvs, T, dt):
    cost_no_ess = 0
    for t in range(T):
        net = loads[t] - pvs[t]
        if net > 0:
            cost_no_ess += prices[t] * net * dt
    return cost_no_ess
def main():
    data_file = '附件1.xlsx'
    times_str, prices, loads, pvs, T, dt = load_data(data_file)
    prob, P_buy, P_ch, P_dis, E = solve_microgrid(prices, loads, pvs, T, dt)
    verify_results(P_buy, P_ch, P_dis, E, prices, loads, pvs, T, dt)
    tables = generate_tables(times_str, prices, loads, pvs, P_buy, P_ch, P_dis, E, T, dt)
    save_to_excel(times_str, prices, loads, pvs, P_buy, P_ch, P_dis, E, T, dt, tables)
    cost_no_ess = calc_no_ess_cost(prices, loads, pvs, T, dt)
    cost_with_ess = tables['total_cost']
    saving = cost_no_ess - cost_with_ess
    saving_pct = saving / cost_no_ess * 100
    print(f"未调整策略前的总价：{cost_no_ess:.2f} 元")
    print(f"调整后的总价：     {cost_with_ess:.2f} 元")
    print(f"下降的价格：      {saving:.2f} 元")
    print(f"下降的百分比：     {saving_pct:.2f}%")
if __name__ == "__main__":
    main()
